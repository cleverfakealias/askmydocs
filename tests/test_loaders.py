from pathlib import Path

import pytest
from docx import Document as DocxDocument

from palimpsest.documents.loaders import UnsupportedFileError, load_file


def test_text_file_becomes_one_document(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("Plain notes about the project.", encoding="utf-8")

    docs = load_file(path)

    assert len(docs) == 1
    assert docs[0].page_content == "Plain notes about the project."
    assert docs[0].metadata == {"source": "notes.txt", "file_type": ".txt"}


def test_markdown_keeps_headings(tmp_path: Path) -> None:
    path = tmp_path / "guide.MD"
    path.write_text("# Setup\n\nInstall the package.", encoding="utf-8")

    docs = load_file(path)

    assert docs[0].page_content.startswith("# Setup")
    assert docs[0].metadata["file_type"] == ".md"


def test_docx_paragraphs_are_extracted(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    document = DocxDocument()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph.")
    document.save(str(path))

    docs = load_file(path)

    assert docs[0].page_content == "First paragraph.\nSecond paragraph."


def test_empty_file_produces_no_documents(tmp_path: Path) -> None:
    path = tmp_path / "blank.txt"
    path.write_text("   \n", encoding="utf-8")

    assert load_file(path) == []


def test_unsupported_extension_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG")

    with pytest.raises(UnsupportedFileError, match=r"unsupported type '\.png'"):
        load_file(path)
